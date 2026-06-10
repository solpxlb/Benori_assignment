"""File exports for raw data, processed data, clusters, and FMCG DealBrief.

Each export writes to disk and returns bytes for Streamlit download buttons.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document


@dataclass
class ExportArtifact:
    """One generated export file: disk path plus in-memory bytes."""

    path: Path
    data: bytes
    filename: str
    mime: str


def _ensure_dir(output_dir: Path | str) -> Path:
    """Create and return the export output directory."""
    path = Path(output_dir)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _json_safe(value: Any) -> Any:
    """Convert pandas/numpy/datetime values into JSON-safe plain objects."""
    if value is None:
        return None
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_json_safe(v) for v in value]
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    if hasattr(value, "isoformat") and not isinstance(value, str):
        try:
            return value.isoformat()
        except TypeError:
            pass
    if hasattr(value, "item"):
        try:
            return value.item()
        except (TypeError, ValueError):
            pass
    return value


def _df_records(df: pd.DataFrame) -> list[dict]:
    """DataFrame records converted to JSON-safe values."""
    return [_json_safe(record) for record in df.to_dict(orient="records")]


def _write_artifact(output_dir: Path, filename: str, data: bytes, mime: str) -> ExportArtifact:
    """Write bytes to disk and return an ExportArtifact."""
    path = output_dir / filename
    path.write_bytes(data)
    return ExportArtifact(path=path, data=data, filename=filename, mime=mime)


def export_dataframe_csv(df: pd.DataFrame, output_dir: Path | str, filename: str) -> ExportArtifact:
    """Export a DataFrame as CSV."""
    out = _ensure_dir(output_dir)
    data = df.to_csv(index=False).encode("utf-8")
    return _write_artifact(out, filename, data, "text/csv")


def export_dataframe_json(df: pd.DataFrame, output_dir: Path | str, filename: str) -> ExportArtifact:
    """Export a DataFrame as pretty JSON records."""
    out = _ensure_dir(output_dir)
    data = json.dumps(_df_records(df), indent=2, ensure_ascii=False).encode("utf-8")
    return _write_artifact(out, filename, data, "application/json")


def export_clusters_json(clusters: list[dict], output_dir: Path | str,
                         filename: str = "deal_clusters.json") -> ExportArtifact:
    """Export deal clusters as pretty JSON."""
    out = _ensure_dir(output_dir)
    data = json.dumps(_json_safe(clusters), indent=2, ensure_ascii=False).encode("utf-8")
    return _write_artifact(out, filename, data, "application/json")


def export_newsletter_docx(newsletter: dict, output_dir: Path | str,
                           filename: str = "newsletter.docx") -> ExportArtifact:
    """Export the structured newsletter as a Word-compatible DOCX."""
    out = _ensure_dir(output_dir)
    doc = Document()
    doc.add_heading(newsletter["title"], level=0)
    doc.add_paragraph(
        f"{newsletter['data_mode']} | {newsletter['date_range']} | Run: {newsletter['run_timestamp']}"
    )

    snapshot = newsletter["snapshot"]
    doc.add_heading("Executive Snapshot", level=1)
    for text in [
        f"{snapshot['deal_events']} deal event(s) from {snapshot['article_count']} article(s) across {snapshot['source_count']} source(s).",
        f"Confidence mix: {snapshot['confidence_counts'] or {'Low': 0}}.",
        f"Top categories: {', '.join(snapshot['top_categories']) if snapshot['top_categories'] else 'None'}.",
        snapshot["auto_summary"],
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("Top Deal Highlights", level=1)
    _add_cluster_section(doc, newsletter["highlights"], fallback="No high- or medium-confidence deal clusters met the current thresholds.")

    doc.add_heading("Watchlist", level=1)
    _add_watchlist_section(doc, newsletter["watchlist"], fallback="No low-confidence/watchlist clusters.")

    doc.add_heading("Methodology", level=1)
    for sentence in newsletter["methodology"]:
        doc.add_paragraph(sentence, style="List Bullet")

    doc.add_heading("Source Appendix", level=1)
    if newsletter["source_appendix"]:
        for source in newsletter["source_appendix"]:
            doc.add_paragraph(
                f"{source['cluster_id']} | {source['label']} | {source['title']} | {source['url']}",
                style="List Bullet",
            )
    else:
        doc.add_paragraph("No cited URLs.")

    buf = BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    return _write_artifact(out, filename, data,
                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def _add_cluster_section(doc: Document, clusters: list[dict], fallback: str) -> None:
    """Append cluster summaries to a DOCX document."""
    if not clusters:
        doc.add_paragraph(fallback)
        return
    for cluster in clusters:
        p = doc.add_paragraph()
        p.add_run(cluster.get("canonical_headline", "Untitled deal event")).bold = True
        companies = cluster.get("companies")
        companies_text = ", ".join(companies) if isinstance(companies, list) else (companies or "Not clearly identified")
        doc.add_paragraph(
            f"{cluster.get('deal_type', 'unknown')} | {companies_text} | "
            f"{cluster.get('geography', 'Not specified')} | {cluster.get('deal_value', 'undisclosed')} | "
            f"{cluster.get('confidence', 'Low')} confidence | "
            f"rel {cluster.get('relevance_score', 0)} / cred {cluster.get('credibility_score', 0)} | "
            f"{cluster.get('source_count', 0)} source(s)"
        )
        doc.add_paragraph(cluster.get("why_it_matters", ""))
        for source in cluster.get("sources", []):
            doc.add_paragraph(
                f"{source.get('source_name') or source.get('domain') or 'source'}: "
                f"{source.get('title', '')} {source.get('url', '')}",
                style="List Bullet",
            )


def _add_watchlist_section(doc: Document, clusters: list[dict], fallback: str) -> None:
    """Append compact watchlist one-liners to a DOCX document."""
    if not clusters:
        doc.add_paragraph(fallback)
        return
    for cluster in clusters:
        companies = cluster.get("companies")
        companies_text = ", ".join(companies) if isinstance(companies, list) else (companies or "Not clearly identified")
        doc.add_paragraph(
            f"{cluster.get('canonical_headline', 'Untitled deal event')} | "
            f"{cluster.get('deal_type', 'unknown')} | {companies_text} | "
            f"{cluster.get('deal_value', 'undisclosed')} | {cluster.get('confidence', 'Low')} | "
            f"rel {cluster.get('relevance_score', 0)} / cred {cluster.get('credibility_score', 0)} | "
            f"{cluster.get('source_count', 0)} source(s)",
            style="List Bullet",
        )


def _cell_safe(value: Any) -> Any:
    """Make a value safe for Excel cells."""
    safe = _json_safe(value)
    if isinstance(safe, (dict, list)):
        return json.dumps(safe, ensure_ascii=False)
    return "" if safe is None else safe


def _sheet_df(records: list[dict], columns: list[str]) -> pd.DataFrame:
    """Build a sheet DataFrame with stable columns."""
    if not records:
        return pd.DataFrame(columns=columns)
    df = pd.DataFrame(records)
    for col in columns:
        if col not in df.columns:
            df[col] = ""
    return df[columns].map(_cell_safe)


def _clusters_flat(clusters: list[dict]) -> pd.DataFrame:
    """Flatten cluster records for an Excel sheet."""
    rows = []
    for cluster in clusters:
        row = {k: v for k, v in cluster.items() if k != "sources"}
        row["companies"] = ", ".join(v for v in cluster.get("companies", []) if v) if isinstance(cluster.get("companies"), list) else cluster.get("companies", "")
        row["source_titles"] = "; ".join(s.get("title", "") for s in cluster.get("sources", []))
        rows.append(row)
    columns = [
        "cluster_id", "canonical_headline", "deal_type", "companies", "geography",
        "category", "deal_value", "relevance_score", "credibility_score",
        "confidence", "why_it_matters", "article_count", "source_count",
        "earliest_published_at", "latest_published_at", "is_sample", "source_titles",
    ]
    return _sheet_df(rows, columns)


def _evidence_flat(clusters: list[dict]) -> pd.DataFrame:
    """Flatten cluster source evidence for Excel."""
    rows = []
    for cluster in clusters:
        for source in cluster.get("sources", []):
            rows.append({
                "cluster_id": cluster.get("cluster_id", ""),
                "canonical_headline": cluster.get("canonical_headline", ""),
                "deal_type": cluster.get("deal_type", ""),
                "source_name": source.get("source_name", ""),
                "domain": source.get("domain", ""),
                "credibility_tier": source.get("credibility_tier", ""),
                "published_at": source.get("published_at", ""),
                "title": source.get("title", ""),
                "url": source.get("url", ""),
            })
    return _sheet_df(rows, [
        "cluster_id", "canonical_headline", "deal_type", "source_name", "domain",
        "credibility_tier", "published_at", "title", "url",
    ])


def _newsletter_rows(clusters: list[dict]) -> list[dict]:
    """Rows for highlight/watchlist sheets."""
    rows = []
    for cluster in clusters:
        companies = cluster.get("companies")
        rows.append({
            "cluster_id": cluster.get("cluster_id", ""),
            "headline": cluster.get("canonical_headline", ""),
            "deal_type": cluster.get("deal_type", ""),
            "companies": ", ".join(companies) if isinstance(companies, list) else companies,
            "category": cluster.get("category", ""),
            "geography": cluster.get("geography", ""),
            "deal_value": cluster.get("deal_value", ""),
            "confidence": cluster.get("confidence", ""),
            "relevance_score": cluster.get("relevance_score", ""),
            "credibility_score": cluster.get("credibility_score", ""),
            "source_count": cluster.get("source_count", ""),
            "why_it_matters": cluster.get("why_it_matters", ""),
        })
    return rows


def _rejected_duplicates(processed_articles: pd.DataFrame) -> pd.DataFrame:
    """Rows rejected or marked duplicate for transparency export."""
    if processed_articles.empty:
        return processed_articles.copy()
    mask = (
        (processed_articles.get("status", pd.Series("", index=processed_articles.index)) == "reject")
        | (processed_articles.get("dedupe_status", pd.Series("", index=processed_articles.index)) == "duplicate")
    )
    return processed_articles.loc[mask].copy()


def export_newsletter_xlsx(newsletter: dict, clusters: list[dict], processed_articles: pd.DataFrame,
                           output_dir: Path | str, filename: str = "newsletter.xlsx") -> ExportArtifact:
    """Export the newsletter workbook with seven reviewer-friendly sheets."""
    out = _ensure_dir(output_dir)
    buf = BytesIO()

    snapshot = newsletter["snapshot"]
    snapshot_df = _sheet_df([
        {"metric": "Data mode", "value": newsletter["data_mode"]},
        {"metric": "Date range", "value": newsletter["date_range"]},
        {"metric": "Run timestamp", "value": newsletter["run_timestamp"]},
        {"metric": "Deal events", "value": snapshot["deal_events"]},
        {"metric": "Article evidence rows", "value": snapshot["article_count"]},
        {"metric": "Unique sources", "value": snapshot["source_count"]},
        {"metric": "Confidence mix", "value": snapshot["confidence_counts"]},
        {"metric": "Top categories", "value": ", ".join(snapshot["top_categories"])},
        {"metric": "Auto summary", "value": snapshot["auto_summary"]},
    ], ["metric", "value"])

    highlight_df = _sheet_df(_newsletter_rows(newsletter["highlights"]), [
        "cluster_id", "headline", "deal_type", "companies", "category", "geography",
        "deal_value", "confidence", "relevance_score", "credibility_score",
        "source_count", "why_it_matters",
    ])
    watchlist_df = _sheet_df(_newsletter_rows(newsletter["watchlist"]), list(highlight_df.columns))
    methodology_df = _sheet_df(
        [{"section": "Methodology", "text": s} for s in newsletter["methodology"]]
        + [{"section": "Source Appendix", "text": f"{s['cluster_id']} | {s['label']} | {s['title']} | {s['url']}"} for s in newsletter["source_appendix"]],
        ["section", "text"],
    )

    sheets = {
        "Executive Snapshot": snapshot_df,
        "Deal Highlights": highlight_df,
        "Deal Clusters": _clusters_flat(clusters),
        "Article Evidence": _evidence_flat(clusters),
        "Watchlist": watchlist_df,
        "Rejected & Duplicates": _rejected_duplicates(processed_articles).map(_cell_safe),
        "Methodology": methodology_df,
    }

    with pd.ExcelWriter(buf, engine="xlsxwriter") as writer:
        for sheet_name, df in sheets.items():
            _write_sheet(writer, df, sheet_name)

    data = buf.getvalue()
    return _write_artifact(out, filename, data,
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _write_sheet(writer: pd.ExcelWriter, df: pd.DataFrame, sheet_name: str) -> None:
    """Write one formatted sheet to an XlsxWriter workbook."""
    safe_df = df.copy()
    if safe_df.empty and len(safe_df.columns) == 0:
        safe_df = pd.DataFrame({"note": []})
    safe_df.to_excel(writer, sheet_name=sheet_name, index=False)
    workbook = writer.book
    worksheet = writer.sheets[sheet_name]
    header_fmt = workbook.add_format({"bold": True, "bg_color": "#D9EAF7", "border": 1})
    for col_num, value in enumerate(safe_df.columns):
        worksheet.write(0, col_num, value, header_fmt)
        width = max(12, min(45, max([len(str(value))] + [len(str(v)) for v in safe_df[value].head(50)])))
        worksheet.set_column(col_num, col_num, width)
    worksheet.freeze_panes(1, 0)
    if len(safe_df.columns) > 0:
        worksheet.autofilter(0, 0, max(len(safe_df), 1), len(safe_df.columns) - 1)


def export_all(
    raw_articles: pd.DataFrame,
    processed_articles: pd.DataFrame,
    clusters: list[dict],
    newsletter_markdown: str,
    newsletter: dict,
    output_dir: Path | str = "outputs",
) -> dict[str, ExportArtifact]:
    """Write every Phase 7 export and return artifacts keyed by logical name."""
    out = _ensure_dir(output_dir)
    artifacts = {
        "raw_csv": export_dataframe_csv(raw_articles, out, "raw_articles.csv"),
        "raw_json": export_dataframe_json(raw_articles, out, "raw_articles.json"),
        "processed_csv": export_dataframe_csv(processed_articles, out, "processed_articles.csv"),
        "processed_json": export_dataframe_json(processed_articles, out, "processed_articles.json"),
        "clusters_json": export_clusters_json(clusters, out, "deal_clusters.json"),
        "newsletter_docx": export_newsletter_docx(newsletter, out, "newsletter.docx"),
        "newsletter_xlsx": export_newsletter_xlsx(newsletter, clusters, processed_articles, out, "newsletter.xlsx"),
    }
    artifacts["newsletter_markdown"] = _write_artifact(
        out, "newsletter.md", newsletter_markdown.encode("utf-8"), "text/markdown"
    )
    return artifacts
